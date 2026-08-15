#!/usr/bin/env python3
"""Local server and SQLite persistence for TradeFlow Atlas."""
from __future__ import annotations

import base64
import binascii
import json
import os
import re
import sqlite3
import ssl
import subprocess
import tempfile
import threading
from datetime import date, timedelta
from html import unescape
from http import HTTPStatus
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from io import BytesIO
from pathlib import Path
from urllib.parse import parse_qs, parse_qsl, unquote, urlencode, urlparse, urlunparse
from urllib.error import HTTPError, URLError
from urllib.request import Request, urlopen

ROOT = Path(__file__).resolve().parent
# Defaults to the local outputs directory. On Railway, mount a Volume at /data
# and set TRADEFLOW_DATA_DIR=/data so notes and progress survive redeploys.
DATA_DIR = Path(os.environ.get("TRADEFLOW_DATA_DIR", ROOT))
DATA_DIR.mkdir(parents=True, exist_ok=True)
DATABASE = DATA_DIR / "tradeflow-atlas.sqlite3"
CURRICULUM = ROOT.parent / "work" / "curriculum-source" / "curriculum"
EXPORTER = ROOT.parent / "work" / "export_tradeflow_xlsx.mjs"
LOCAL_NODE = os.environ.get("NODE_BINARY", "/Users/abhyudayamarya/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node")
ENV_FILE = ROOT / ".env"
SIGNAL_CACHE: dict[str, object] = {"created_at": 0.0, "items": []}
NEWS_MODEL_STATE: dict[str, object] = {"model": None, "topic_embeddings": None, "unavailable": False}
NEWS_MODEL_LOCK = threading.Lock()
IAM_TOKEN_CACHE: dict[str, object] = {"token": "", "expires_at": 0.0}
MAX_REQUEST_BYTES = 28 * 1024 * 1024
MAX_ATTACHMENTS = 4
MAX_ATTACHMENT_BYTES = 10 * 1024 * 1024
MAX_TOTAL_ATTACHMENT_BYTES = 20 * 1024 * 1024
MAX_DOCUMENT_CHARS = 60_000
MAX_PDF_PAGES = 50
MAX_RENDERED_PAGES = 6
MAX_HISTORY_MESSAGES = 12
MAX_HISTORY_MESSAGE_CHARS = 4_000

NEWS_TOPICS: dict[str, tuple[str, ...]] = {
    "Demand planning": (
        "demand forecasting", "demand sensing", "forecast accuracy", "consumer demand", "demand planning",
    ),
    "Supply planning": (
        "supply planning", "capacity planning", "production planning", "supply constraints", "scenario planning",
    ),
    "Inventory": (
        "inventory optimisation", "inventory reduction", "safety stock", "stock availability", "retailer destocking",
    ),
    "Retail & CPG": (
        "retail operations", "consumer packaged goods", "FMCG", "retail planning", "merchandise planning",
    ),
    "TPM & NRM": (
        "trade promotion management", "trade spend", "promotion effectiveness", "promotional uplift", "net revenue management",
    ),
    "Logistics": (
        "transportation", "freight", "warehousing", "shipping", "distribution network", "last mile logistics",
    ),
    "S&OP / IBP": (
        "integrated business planning", "sales and operations planning", "S&OP", "IBP", "cross-functional planning",
    ),
    "Procurement": (
        "procurement", "strategic sourcing", "supplier management", "supplier risk", "purchasing",
    ),
    "Supply-chain AI": (
        "artificial intelligence in supply chain", "generative AI for retail", "AI demand forecasting", "supply chain copilot", "autonomous planning",
    ),
}
NEWS_TOPIC_SENTENCES: tuple[tuple[str, str], ...] = tuple(
    (category, f"An industry article about {phrase} in supply-chain, retail, or consumer-goods operations.")
    for category, phrases in NEWS_TOPICS.items()
    for phrase in phrases
)
NEWS_BLOCKED_HOSTS = {"biztoc.com"}
NEWS_TRACKING_PARAMETERS = {"fbclid", "gclid", "mc_cid", "mc_eid", "ref", "referrer"}


def load_local_env() -> None:
    """Load only this app's local .env file; explicitly exported shell values win."""
    if not ENV_FILE.is_file():
        return
    for raw_line in ENV_FILE.read_text(encoding="utf-8", errors="ignore").splitlines():
        match = re.match(r"\s*(?:export\s+)?([A-Za-z_][A-Za-z0-9_]*)\s*=\s*(.*?)\s*$", raw_line)
        if not match:
            continue
        key, value = match.groups()
        if (value.startswith('"') and value.endswith('"')) or (value.startswith("'") and value.endswith("'")):
            value = value[1:-1]
        else:
            value = value.split(" #", 1)[0].strip()
        if value and not os.environ.get(key):
            os.environ[key] = value


load_local_env()


def connection() -> sqlite3.Connection:
    db = sqlite3.connect(DATABASE)
    db.row_factory = sqlite3.Row
    db.execute("CREATE TABLE IF NOT EXISTS notes (zone TEXT PRIMARY KEY, content TEXT NOT NULL, updated_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    db.execute("CREATE TABLE IF NOT EXISTS note_entries (id INTEGER PRIMARY KEY AUTOINCREMENT, zone TEXT NOT NULL, content TEXT NOT NULL, created_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    db.execute("CREATE TABLE IF NOT EXISTS capsules (zone TEXT PRIMARY KEY, content TEXT NOT NULL, unlock_date TEXT NOT NULL, updated_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    db.execute("CREATE TABLE IF NOT EXISTS progress (zone TEXT PRIMARY KEY, checked_keys TEXT NOT NULL, updated_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    db.execute("CREATE TABLE IF NOT EXISTS settings (key TEXT PRIMARY KEY, value TEXT NOT NULL, updated_at TEXT DEFAULT CURRENT_TIMESTAMP)")
    columns = {row[1] for row in db.execute("PRAGMA table_info(progress)")}
    if "checked_keys" not in columns:
        db.execute("ALTER TABLE progress ADD COLUMN checked_keys TEXT NOT NULL DEFAULT '[]'")
    db.execute("INSERT INTO note_entries(zone, content, created_at) SELECT zone, content, updated_at FROM notes WHERE NOT EXISTS (SELECT 1 FROM note_entries n WHERE n.zone = notes.zone AND n.content = notes.content)")
    return db


def plain_note(value: str) -> str:
    """Keep exports readable when rich-text notes contain browser HTML."""
    value = re.sub(r"<(br|/p|/div|/li|/h[1-6])[^>]*>", "\n", value, flags=re.I)
    value = re.sub(r"<[^>]+>", "", value)
    return re.sub(r"\n{3,}", "\n\n", value).strip()


def build_excel_export(payload: dict[str, list[dict[str, object]]]) -> bytes:
    """Create the local export without a machine-specific Node.js dependency."""
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    except ImportError:
        # The existing local Codex runtime can still export before pip installs
        # the Railway dependency. Production always takes the portable path above.
        node = Path(LOCAL_NODE)
        if not node.is_file() or not EXPORTER.is_file():
            raise RuntimeError("Excel export needs openpyxl. Run: pip install -r requirements.txt")
        with tempfile.TemporaryDirectory() as temp_dir:
            data_file = Path(temp_dir) / "tradeflow-data.json"
            output_file = Path(temp_dir) / "tradeflow-atlas-export.xlsx"
            data_file.write_text(json.dumps(payload), encoding="utf-8")
            subprocess.run([str(node), str(EXPORTER), str(data_file), str(output_file)], check=True, capture_output=True, text=True)
            return output_file.read_bytes()

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = (
        ("Rough Notes", ["Learning zone", "Note", "Last updated"], payload["notes"], ("zone", "content", "updated_at")),
        ("Time Capsules", ["Learning zone", "Message", "Unlock date", "Last updated"], payload["capsules"], ("zone", "content", "unlock_date", "updated_at")),
        ("Checklist Progress", ["Learning zone", "Completed checklist keys", "Last updated"], payload["progress"], ("zone", "checked_keys", "updated_at")),
        ("Settings", ["Setting", "Value", "Last updated"], payload["settings"], ("key", "value", "updated_at")),
    )
    border = Border(left=Side(style="thin", color="E3E1F0"), right=Side(style="thin", color="E3E1F0"), top=Side(style="thin", color="E3E1F0"), bottom=Side(style="thin", color="E3E1F0"))
    for sheet_name, headers, rows, fields in sheets:
        sheet = workbook.create_sheet(sheet_name)
        sheet.append(headers)
        for row in rows or [{}]:
            sheet.append([row.get(field, "") if row else "No saved entries yet" if index == 0 else "" for index, field in enumerate(fields)])
        for cell in sheet[1]:
            cell.fill = PatternFill("solid", fgColor="24176C")
            cell.font = Font(bold=True, color="FFFFFF")
            cell.alignment = Alignment(horizontal="center")
            cell.border = border
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
                cell.border = border
        sheet.freeze_panes = "A2"
        sheet.column_dimensions["A"].width = 24
        for column in range(2, len(headers) + 1):
            sheet.column_dimensions[chr(64 + column)].width = 34
    result = BytesIO()
    workbook.save(result)
    return result.getvalue()


def progress_text(value: str) -> str:
    try:
        entries = json.loads(value)
    except json.JSONDecodeError:
        return value
    return "\n".join(f"✓ {x.get('label', x.get('key', 'Completed task'))}" if isinstance(x, dict) else f"✓ {x}" for x in entries)


def trusted_urlopen(request: Request, timeout: int | None):
    """Use certifi in the project virtual environment on macOS, where system roots can be unavailable."""
    try:
        import certifi
        context = ssl.create_default_context(cafile=certifi.where())
    except ImportError:
        context = ssl.create_default_context()
    return urlopen(request, timeout=timeout, context=context)


def relevant_curriculum(question: str, limit: int = 3) -> list[dict[str, str]]:
    """Retrieve a small, relevant slice of the local curriculum without notes."""
    tokens = {item.lower() for item in re.findall(r"[A-Za-z][A-Za-z0-9&/-]{2,}", question)}
    candidates: list[tuple[int, str, str]] = []
    for file in CURRICULUM.glob("*.md"):
        text = file.read_text(encoding="utf-8", errors="ignore")
        chunks = re.split(r"\n(?=#{1,3}\s)", text)
        for chunk in chunks:
            cleaned = re.sub(r"\s+", " ", chunk).strip()
            # Code-heavy notebook cells are poor retrieval context for a learning answer.
            if len(cleaned) < 70 or cleaned.startswith("# ===") or re.search(r"\bdef\s+\w+\s*\(|client\.chat|response\.choices", cleaned):
                continue
            score = sum(cleaned.lower().count(token) for token in tokens)
            if score:
                title_match = re.search(r"^#{1,3}\s+(.+)", chunk, re.M)
                title = title_match.group(1).strip("# =-") if title_match else file.stem.replace("-", " ")
                candidates.append((score, title, cleaned[:700]))
    candidates.sort(key=lambda item: item[0], reverse=True)
    return [{"title": title, "excerpt": excerpt, "source": f"Your curriculum · {title}"} for _, title, excerpt in candidates[:limit]]


def notion_relevant_context(question: str, limit: int = 3) -> list[dict[str, str]]:
    """Retrieve a small set of relevant Notion pages for the current question, if a token is configured."""
    token = (os.environ.get("NOTION_API") or "").strip()
    if not token or not question.strip():
        return []
    tokens = {item.lower() for item in re.findall(r"[A-Za-z][A-Za-z0-9&/-]{2,}", question)}
    headers = {
        "Authorization": f"Bearer {token}",
        "Content-Type": "application/json",
        "Notion-Version": "2022-06-28",
    }
    try:
        search_request = Request(
            "https://api.notion.com/v1/search",
            data=json.dumps({
                "query": question,
                "filter": {"property": "object", "value": "page"},
                "page_size": max(5, limit * 5),
            }).encode("utf-8"),
            headers=headers,
            method="POST",
        )
        with trusted_urlopen(search_request, timeout=10) as response:
            payload = json.loads(response.read().decode("utf-8", errors="ignore"))
    except Exception:
        return []

    results: list[tuple[int, str, str, str]] = []
    for page in payload.get("results", [])[: max(5, limit * 5)]:
        if page.get("object") != "page":
            continue
        title = ""
        properties = page.get("properties") or {}
        for key, value in properties.items():
            if not isinstance(value, dict):
                continue
            if value.get("type") == "title":
                title = "".join(item.get("plain_text", "") for item in value.get("title", []) if isinstance(item, dict))
                break
            if key == "title" and isinstance(value.get("title"), list):
                title = "".join(item.get("plain_text", "") for item in value.get("title", []) if isinstance(item, dict))
                break
        if not title:
            title = "Notion page"
        page_id = str(page.get("id") or "")
        try:
            blocks_request = Request(
                f"https://api.notion.com/v1/blocks/{page_id}/children?page_size=100",
                headers=headers,
                method="GET",
            )
            with trusted_urlopen(blocks_request, timeout=10) as block_response:
                block_payload = json.loads(block_response.read().decode("utf-8", errors="ignore"))
        except Exception:
            block_payload = {"results": []}

        block_chunks: list[str] = []
        for block in block_payload.get("results", []):
            if not isinstance(block, dict):
                continue
            block_type = block.get("type")
            content = block.get(block_type) or {}
            rich_text = content.get("rich_text", []) if isinstance(content, dict) else []
            texts: list[str] = []
            for item in rich_text:
                if not isinstance(item, dict):
                    continue
                text_value = item.get("plain_text")
                if text_value:
                    texts.append(str(text_value))
                    continue
                nested = item.get("text")
                if isinstance(nested, dict):
                    text_value = nested.get("content")
                    if text_value:
                        texts.append(str(text_value))
            text = "".join(texts)
            if text:
                block_chunks.append(text)
        excerpt = " ".join(block_chunks)[:700].strip()
        if not excerpt:
            excerpt = title
        title_lower = title.lower()
        excerpt_lower = excerpt.lower()
        score = sum(title_lower.count(token) for token in tokens) + sum(excerpt_lower.count(token) for token in tokens)
        if score or not results:
            results.append((score, title.strip() or "Notion page", excerpt.strip() or title.strip(), f"Notion · {title.strip() or 'Notion page'}"))
    results.sort(key=lambda item: item[0], reverse=True)
    return [{"title": title, "excerpt": excerpt, "source": source} for _, title, excerpt, source in results[:limit]]


def local_guidance(question: str, curriculum: list[dict[str, str]]) -> str:
    """Server-side guidance when live model inference is unavailable."""
    rules = (
        (r"promo|promotion|trade spend|ropi|roi|net revenue", "For promotion economics, start with baseline versus incremental volume, then test trade spend and margin sensitivity. Open Zone 01 for trade-investment foundations and Zone 03 for the Promo ROI Analyzer."),
        (r"claim|deduction|chargeback|debit memo", "Frame the claims question as: reason-code pattern → financial exposure → root cause → owner and corrective action. Use Zone 01 for governance and Zone 03 for the Claims Analysis Agent."),
        (r"forecast|demand|uplift|s&op|ibp|planning", "Separate statistical baseline, promotional uplift, constraints, and the decision needed before asking an AI tool for interpretation. Start in Zone 01 and continue to Zone 02 for planning GenAI patterns."),
        (r"risk|control tower|inventory|supply|service", "Prioritise risk by customer/service impact, financial exposure, time sensitivity, root cause, and escalation owner. Zone 04 contains the Supply Risk Copilot."),
        (r"watson|api|granite|model|rag|agent", "Use the watsonx.ai build standard: API key, regional URL, project ID, source grounding, and functional expert review before recommendations are used."),
        (r"capstone|portfolio|interview|consultant", "Lead the portfolio story with the business problem, analytical method and control points, then the client decision or measurable operating impact. Zone 05 is your consulting portfolio zone."),
    )
    for pattern, answer in rules:
        if re.search(pattern, question, re.I):
            return answer
    if curriculum:
        excerpt = re.sub(r"\s+", " ", curriculum[0]["excerpt"]).strip()
        return f"From your curriculum: {excerpt[:420]}\n\nFor a sharper answer, name the planning area, retailer/CPG context, and decision you are trying to make."
    return "I can help with promotion ROI, claims, S&OP, supply risk, watsonx.ai, or your capstone portfolio. Tell me the decision you are trying to make."


def web_search(question: str, limit: int = 2) -> list[dict[str, str]]:
    """Best-effort public web retrieval. It sends only the question, never saved notes."""
    try:
        url = "https://html.duckduckgo.com/html/?" + urlencode({"q": question})
        request = Request(url, headers={"User-Agent": "TradeFlow-Atlas/1.0 learning research"})
        with trusted_urlopen(request, timeout=3) as response:
            page = response.read().decode("utf-8", errors="ignore")
        matches = re.findall(r'<a[^>]+class="result__a"[^>]+href="([^"]+)"[^>]*>(.*?)</a>', page, re.S)
        results = []
        for href, title in matches:
            clean_title = re.sub(r"<[^>]+>", "", unescape(title)).strip()
            parsed = urlparse(unescape(href))
            query = parse_qs(parsed.query)
            link = unquote(query.get("uddg", [href])[0])
            source_host = (urlparse(link).hostname or "").lower()
            if clean_title and link.startswith(("https://", "http://")) and source_host != "biztoc.com" and not source_host.endswith(".biztoc.com"):
                results.append({"title": clean_title[:150], "url": link})
            if len(results) >= limit:
                break
        return results
    except Exception:
        return []


def news_api_key() -> str:
    """Allow the simple NEWS_API name used in this project's .env file."""
    return os.environ.get("NEWS_API") or os.environ.get("NEWS_API_KEY") or os.environ.get("NEWSAPI_KEY", "")


def canonical_news_url(value: str) -> str:
    """Remove fragments and common tracking parameters before URL deduplication."""
    parsed = urlparse(value)
    clean_query = [
        (key, item)
        for key, item in parse_qsl(parsed.query, keep_blank_values=False)
        if not key.lower().startswith("utm_") and key.lower() not in NEWS_TRACKING_PARAMETERS
    ]
    return urlunparse((parsed.scheme.lower(), parsed.netloc.lower(), parsed.path.rstrip("/"), "", urlencode(clean_query), ""))


def news_semantic_model() -> tuple[object | None, object | None]:
    """Load the small retrieval model once; keep news available if it cannot load."""
    if NEWS_MODEL_STATE["unavailable"]:
        return None, None
    if NEWS_MODEL_STATE["model"] is not None:
        return NEWS_MODEL_STATE["model"], NEWS_MODEL_STATE["topic_embeddings"]
    with NEWS_MODEL_LOCK:
        if NEWS_MODEL_STATE["model"] is not None:
            return NEWS_MODEL_STATE["model"], NEWS_MODEL_STATE["topic_embeddings"]
        try:
            from sentence_transformers import SentenceTransformer

            model_name = os.environ.get("NEWS_EMBEDDING_MODEL_ID", "sentence-transformers/all-MiniLM-L6-v2")
            model = SentenceTransformer(model_name)
            topic_embeddings = model.encode(
                [sentence for _, sentence in NEWS_TOPIC_SENTENCES],
                normalize_embeddings=True,
                show_progress_bar=False,
            )
            NEWS_MODEL_STATE.update({"model": model, "topic_embeddings": topic_embeddings})
            return model, topic_embeddings
        except Exception:
            NEWS_MODEL_STATE["unavailable"] = True
            return None, None


def lexical_news_score(title: str, description: str) -> tuple[float, str]:
    """Taxonomy-aware fallback for machines that cannot load Sentence Transformers."""
    title_text = title.lower()
    body_text = f"{title} {description}".lower()
    best_category = "Supply chain"
    best_score = 0.0
    for category, phrases in NEWS_TOPICS.items():
        title_hits = sum(1 for phrase in phrases if phrase.lower() in title_text)
        body_hits = sum(1 for phrase in phrases if phrase.lower() in body_text)
        score = min(1.0, title_hits * 0.42 + max(0, body_hits - title_hits) * 0.19)
        if score > best_score:
            best_score, best_category = score, category
    # Broad discovery terms are deliberately less specific than the taxonomy
    # phrases above. They still need a useful category and a non-zero score when
    # Sentence Transformers is unavailable (for example on a fresh Railway
    # deploy while the embedding model is warming up).
    broad_categories = {
        "Supply planning": ("supply chain", "supply network", "supply disruption"),
        "Demand planning": ("demand forecast", "forecasting", "consumer demand"),
        "Inventory": ("inventory", "stockout", "stock level"),
        "Retail & CPG": ("retail", "cpg", "fmcg", "consumer goods"),
        "TPM & NRM": ("promotion", "trade spend", "revenue management"),
        "Logistics": ("logistics", "freight", "shipping", "warehouse"),
        "Procurement": ("procurement", "sourcing", "supplier"),
        "Supply-chain AI": ("artificial intelligence", "generative ai", "copilot"),
    }
    for category, terms in broad_categories.items():
        title_hits = sum(1 for term in terms if term in title_text)
        body_hits = sum(1 for term in terms if term in body_text)
        broad_score = min(0.64, title_hits * 0.30 + max(0, body_hits - title_hits) * 0.14)
        if broad_score > best_score:
            best_score, best_category = broad_score, category
    broad_hits = sum(
        term in body_text
        for term in ("supply chain", "logistics", "retail", "consumer goods", "inventory", "planning", "promotion", "procurement")
    )
    if best_score and broad_hits:
        best_score = min(1.0, best_score + min(0.18, broad_hits * 0.04))
    return best_score, best_category


def rank_industry_candidates(candidates: list[dict[str, str]]) -> tuple[list[dict[str, object]], str]:
    """Rank, label, and semantically deduplicate recent industry coverage."""
    if not candidates:
        return [], "taxonomy"
    semantic_minimum = float(os.environ.get("NEWS_RELEVANCE_THRESHOLD", "0.30"))
    model, topic_embeddings = news_semantic_model()
    ranked: list[dict[str, object]] = []
    article_embeddings = None
    if model is not None and topic_embeddings is not None:
        try:
            article_embeddings = model.encode(
                [f"{item['title']}. {item['description']}" for item in candidates],
                normalize_embeddings=True,
                show_progress_bar=False,
            )
            similarities = article_embeddings @ topic_embeddings.T
            for index, item in enumerate(candidates):
                topic_index = int(similarities[index].argmax())
                semantic_score = float(similarities[index][topic_index])
                lexical_score, lexical_category = lexical_news_score(item["title"], item["description"])
                # Treat embeddings and explicit domain language as complementary
                # signals. This prevents an unusually conservative embedding
                # score from hiding a headline that directly names supply-chain,
                # retail, planning, logistics, or TPM concepts.
                if lexical_score > semantic_score:
                    relevance = lexical_score
                    category = lexical_category
                else:
                    relevance = semantic_score
                    category = NEWS_TOPIC_SENTENCES[topic_index][0]
                ranked.append({
                    **item,
                    "relevance": round(relevance, 3),
                    "category": category,
                    "_embedding_index": index,
                })
            method = "semantic + taxonomy"
        except Exception:
            article_embeddings = None
            ranked = []
            method = "taxonomy"
    else:
        method = "taxonomy"
    if not ranked:
        for item in candidates:
            score, category = lexical_news_score(item["title"], item["description"])
            ranked.append({**item, "relevance": round(score, 3), "category": category})
    ranked.sort(key=lambda item: (float(item["relevance"]), item.get("publishedAt", "")), reverse=True)
    # Embedding similarities vary by model and by the wording used in a day's
    # headlines. Keep the configured ceiling, but calibrate the live cutoff to
    # the strongest result so a relevant, sparse news day does not render an
    # empty reading desk. The NewsAPI discovery queries remain domain-scoped.
    strongest = float(ranked[0]["relevance"]) if ranked else 0.0
    if method.startswith("semantic"):
        minimum = min(semantic_minimum, max(0.16, strongest * 0.72))
    else:
        configured_lexical = float(os.environ.get("NEWS_LEXICAL_THRESHOLD", "0.18"))
        minimum = min(configured_lexical, max(0.08, strongest * 0.55))
    ranked = [item for item in ranked if float(item["relevance"]) >= minimum]

    selected: list[dict[str, object]] = []
    selected_title_tokens: list[set[str]] = []
    selected_embedding_indexes: list[int] = []
    for item in ranked:
        title_tokens = set(re.findall(r"[a-z0-9]+", str(item["title"]).lower())) - {"the", "a", "an", "and", "of", "to", "in", "for"}
        lexical_duplicate = any(
            len(title_tokens & previous) / max(1, len(title_tokens | previous)) >= 0.72
            for previous in selected_title_tokens
        )
        semantic_duplicate = False
        embedding_index = item.get("_embedding_index")
        if article_embeddings is not None and isinstance(embedding_index, int):
            semantic_duplicate = any(
                float(article_embeddings[embedding_index] @ article_embeddings[previous]) >= 0.88
                for previous in selected_embedding_indexes
            )
        if lexical_duplicate or semantic_duplicate:
            continue
        item.pop("_embedding_index", None)
        selected.append(item)
        selected_title_tokens.append(title_tokens)
        if isinstance(embedding_index, int):
            selected_embedding_indexes.append(embedding_index)
        if len(selected) >= 12:
            break
    return selected, method


def industry_signals(force: bool = False) -> dict[str, object]:
    """Fetch and semantically rank public supply-chain news; no private notes leave the device."""
    import time

    cached = SIGNAL_CACHE.get("items", [])
    if cached and not force and time.time() - float(SIGNAL_CACHE["created_at"]) < 1800:
        return {
            "items": cached,
            "cached": True,
            "candidateCount": SIGNAL_CACHE.get("candidate_count", len(cached)),
            "ranking": SIGNAL_CACHE.get("ranking", "taxonomy"),
        }
    api_key = news_api_key()
    if not api_key:
        return {"items": [], "error": "Add NEWS_API to outputs/.env to enable Industry Signals."}
    queries = (
        '("supply chain" OR logistics OR inventory OR procurement OR sourcing)',
        '("demand planning" OR "supply planning" OR "sales and operations planning" OR "integrated business planning")',
        '("trade promotion" OR "trade spend" OR "revenue growth management" OR (retail AND (CPG OR FMCG OR "consumer goods")))',
    )
    today = date.today()
    from_date = today - timedelta(days=7)
    try:
        articles = []
        for query in queries:
            url = "https://newsapi.org/v2/everything?" + urlencode({
                "q": query,
                "language": "en",
                "sortBy": "publishedAt",
                "pageSize": "100",
                "from": from_date.isoformat(),
                "to": today.isoformat(),
            })
            request = Request(url, headers={"X-Api-Key": api_key, "User-Agent": "TradeFlow-Atlas/1.0"})
            with trusted_urlopen(request, timeout=12) as response:
                articles.extend(json.loads(response.read().decode("utf-8")).get("articles", []))
        candidates: list[dict[str, str]] = []
        seen_urls: set[str] = set()
        seen_titles: set[str] = set()
        for article in articles:
            title = str(article.get("title") or "").strip()
            description = str(article.get("description") or "").strip()
            article_url = str(article.get("url") or "").strip()
            source_name = str((article.get("source") or {}).get("name") or "Industry source").strip()
            source_host = (urlparse(article_url).hostname or "").lower()
            canonical_url = canonical_news_url(article_url)
            normalized_title = re.sub(r"\W+", "", title.lower())
            if (
                not title
                or not article_url.startswith(("https://", "http://"))
                or not canonical_url
                or canonical_url in seen_urls
                or normalized_title in seen_titles
                or any(source_host == host or source_host.endswith("." + host) for host in NEWS_BLOCKED_HOSTS)
                or "biztoc" in source_name.lower()
                or title == "[Removed]"
            ):
                continue
            seen_urls.add(canonical_url)
            seen_titles.add(normalized_title)
            candidates.append({
                "title": title[:220],
                "description": description[:340],
                "url": canonical_url,
                "source": source_name[:100],
                "publishedAt": str(article.get("publishedAt") or ""),
            })
        items, ranking = rank_industry_candidates(candidates)
        SIGNAL_CACHE.update({
            "created_at": time.time(),
            "items": items,
            "candidate_count": len(candidates),
            "ranking": ranking,
        })
        return {"items": items, "cached": False, "candidateCount": len(candidates), "ranking": ranking}
    except HTTPError as error:
        if cached:
            return {"items": cached, "cached": True, "stale": True, "warning": f"NewsAPI returned HTTP {error.code}; showing the last successful feed."}
        return {"items": [], "error": f"NewsAPI returned HTTP {error.code}. Check the NewsAPI plan, key and request allowance."}
    except URLError as error:
        if cached:
            return {"items": cached, "cached": True, "stale": True, "warning": "NewsAPI is temporarily unreachable; showing the last successful feed."}
        reason = str(getattr(error, "reason", error))[:180]
        return {"items": [], "error": f"The server could not reach NewsAPI: {reason}"}
    except (json.JSONDecodeError, UnicodeDecodeError):
        if cached:
            return {"items": cached, "cached": True, "stale": True, "warning": "NewsAPI returned an unreadable response; showing the last successful feed."}
        return {"items": [], "error": "NewsAPI returned an unreadable response. Try refreshing later."}
    except Exception as error:
        if cached:
            return {"items": cached, "cached": True, "stale": True, "warning": "The feed could not refresh; showing the last successful feed."}
        return {"items": [], "error": f"Industry Signals could not refresh: {type(error).__name__}."}


def watsonx_config() -> tuple[dict[str, str] | None, list[str]]:
    values = {
        "apikey": os.environ.get("WATSONX_APIKEY", ""),
        "project_id": os.environ.get("WATSONX_PROJECT_ID", ""),
        "url": os.environ.get("WATSONX_URL", ""),
        "model_id": os.environ.get("WATSONX_MODEL_ID", "mistralai/mistral-large"),
    }
    missing = [name for name in ("WATSONX_APIKEY", "WATSONX_PROJECT_ID", "WATSONX_URL") if not os.environ.get(name)]
    return (values if not missing else None), missing


def generate_watsonx_answer(question: str, curriculum: list[dict[str, str]], sources: list[dict[str, str]]) -> str:
    config, missing = watsonx_config()
    if not config:
        raise RuntimeError("Configure " + ", ".join(missing) + " locally before enabling live AI.")
    local_context = "\n\n".join(f"[{item['source']}]\n{item['excerpt']}" for item in curriculum) or "No directly matching local excerpt was found."
    web_context = "\n".join(f"- {item['title']}: {item['url']}" for item in sources) or "(None.)"
    system = """You are Anusha, a practical strategic supply-chain planning and retail TPM learning assistant. Use the local curriculum as the primary grounding. Use web sources only as supplemental reading leads; do not claim facts not present in the supplied source list. Do not invent client data, citations, or metrics. Distinguish fact from recommendation. Answer in 120–180 words maximum, using at most 4 short bullets when useful. Finish every sentence and return only the helpful answer."""
    user_content = f"""USER QUESTION:
{question}

LOCAL CURRICULUM (private to this local app; never include saved notes):
{local_context}

WEB RESEARCH LEADS:
{web_context}"""
    import time
    if time.time() >= float(IAM_TOKEN_CACHE["expires_at"]):
        token_request = Request(
            "https://iam.cloud.ibm.com/identity/token",
            data=urlencode({"grant_type": "urn:ibm:params:oauth:grant-type:apikey", "apikey": config["apikey"]}).encode(),
            headers={"Content-Type": "application/x-www-form-urlencoded", "Accept": "application/json"},
        )
        with trusted_urlopen(token_request, timeout=20) as token_response:
            token_payload = json.loads(token_response.read().decode("utf-8"))
        IAM_TOKEN_CACHE.update({"token": token_payload["access_token"], "expires_at": time.time() + int(token_payload.get("expires_in", 3600)) - 60})
    endpoint = config["url"].rstrip("/") + "/ml/v1/text/chat?version=2026-08-02"
    request_body = {
        "model_id": config["model_id"],
        "project_id": config["project_id"],
        "messages": [
            {"role": "system", "content": system},
            {"role": "user", "content": [{"type": "text", "text": user_content}]},
        ],
        "max_tokens": 280,
        "temperature": 0.2,
    }
    
    chat_request = Request(endpoint, data=json.dumps(request_body).encode(), headers={"Authorization": "Bearer " + str(IAM_TOKEN_CACHE["token"]), "Content-Type": "application/json", "Accept": "application/json"})
    try:
        with trusted_urlopen(chat_request, timeout=None) as chat_response:
            payload = json.loads(chat_response.read().decode("utf-8"))
    except HTTPError as error:
        body = error.read().decode("utf-8", errors="ignore")
        raise RuntimeError(f"watsonx.ai HTTP {error.code}: {body}") from error
    
    answer = str(payload["choices"][0]["message"]["content"]).strip()
    # Avoid presenting a half-finished final sentence if a provider stops at its token limit.
    if answer and answer[-1] not in ".!?`":
        complete = re.findall(r".*?[.!?](?:\s|$)", answer, flags=re.S)
        if complete:
            answer = "".join(complete).strip()
    return answer or "I could not form a complete response. Please try a more specific question."


def generate_watsonx_answer_with_timeout(question: str, curriculum: list[dict[str, str]], sources: list[dict[str, str]]) -> str:
    """Compatibility wrapper: inference is intentionally allowed to finish without a time cutoff."""
    return generate_watsonx_answer(question, curriculum, sources)

def huggingface_config() -> tuple[dict[str, str] | None, list[str]]:
    # Keep the legacy HF_MODEL_ID setting working, while allowing Anusha to
    # use specialised Qwen deployments when you add them to .env.
    default_qwen = os.environ.get("HF_MODEL_ID", "Qwen/Qwen3-32B")
    # Keep this default limited to a model Hugging Face documents as being
    # served by Inference Providers. A model can support vision in
    # Transformers without being callable through the hosted router.
    default_vision = "Qwen/Qwen2-VL-72B-Instruct"
    values = {
        "token": os.environ.get("HF_TOKEN", ""),
        "general_model_id": os.environ.get("HF_GENERAL_MODEL_ID", default_qwen),
        "multimodal_model_id": os.environ.get("HF_MULTIMODAL_MODEL_ID", default_vision),
        "multimodal_fallback_model_id": os.environ.get("HF_MULTIMODAL_FALLBACK_MODEL_ID", default_vision),
        "quiz_model_id": os.environ.get("HF_QUIZ_MODEL_ID", default_qwen),
    }
    missing = [name for name in ("HF_TOKEN",) if not os.environ.get(name)]
    return (values if not missing else None), missing


def route_anusha_request(question: str, documents: list[dict[str, object]], mode: str, use_web: bool) -> dict[str, str]:
    """Choose a bounded Qwen capability; this is intentionally not autonomous."""
    if mode == "quiz":
        return {
            "name": "Qwen quiz designer",
            "model_key": "quiz_model_id",
            "purpose": "Creates validated, interactive study questions from the current topic or upload.",
        }
    if documents:
        return {
            "name": "Qwen document analyst",
            "model_key": "multimodal_model_id",
            "purpose": "Reads extracted PDF text and, where needed, attached page or image visuals.",
        }
    if mode == "summarize":
        return {
            "name": "Qwen study summariser",
            "model_key": "general_model_id",
            "purpose": "Produces a concise, curriculum-grounded study summary.",
        }
    if use_web:
        return {
            "name": "Qwen research analyst",
            "model_key": "general_model_id",
            "purpose": "Synthesises current public research leads with the local curriculum.",
        }
    return {
        "name": "Qwen supply-chain practice partner",
        "model_key": "general_model_id",
        "purpose": "Explains SCM, retail TPM and GenAI concepts using the selected curriculum context.",
    }


def prepare_conversation_history(items: object) -> list[dict[str, str]]:
    """Accept a small, text-only conversation window from the current browser session."""
    if items in (None, []):
        return []
    if not isinstance(items, list):
        raise ValueError("Conversation history must be sent as a list.")
    history: list[dict[str, str]] = []
    for item in items[-MAX_HISTORY_MESSAGES:]:
        if not isinstance(item, dict):
            continue
        role = str(item.get("role", "")).strip().lower()
        content = str(item.get("content", "")).strip()
        if role not in {"user", "assistant"} or not content:
            continue
        history.append({"role": role, "content": content[:MAX_HISTORY_MESSAGE_CHARS]})
    return history


def safe_attachment_name(value: object) -> str:
    """Keep uploaded names readable without accepting paths or control characters."""
    name = Path(str(value or "attachment")).name
    name = re.sub(r"[\x00-\x1f\x7f]+", "", name).strip()
    return (name or "attachment")[:160]


def decode_attachment(item: object) -> tuple[str, str, bytes]:
    if not isinstance(item, dict):
        raise ValueError("Each attachment must include a name, type and file data.")
    name = safe_attachment_name(item.get("name"))
    media_type = str(item.get("type") or "").lower().split(";", 1)[0].strip()
    suffix_types = {
        ".pdf": "application/pdf",
        ".png": "image/png",
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".webp": "image/webp",
    }
    media_type = media_type or suffix_types.get(Path(name).suffix.lower(), "")
    allowed = {"application/pdf", "image/png", "image/jpeg", "image/webp"}
    if media_type not in allowed:
        raise ValueError(f"{name}: upload a PDF, PNG, JPEG or WebP file.")
    encoded = str(item.get("data") or "")
    if encoded.startswith("data:"):
        try:
            header, encoded = encoded.split(",", 1)
        except ValueError as error:
            raise ValueError(f"{name}: malformed file data.") from error
        declared = header[5:].split(";", 1)[0].lower()
        if declared and declared != media_type:
            raise ValueError(f"{name}: file type does not match its contents.")
    try:
        raw = base64.b64decode(encoded, validate=True)
    except (binascii.Error, ValueError) as error:
        raise ValueError(f"{name}: could not decode the uploaded file.") from error
    if not raw:
        raise ValueError(f"{name}: the uploaded file is empty.")
    if len(raw) > MAX_ATTACHMENT_BYTES:
        raise ValueError(f"{name}: files must be 10 MB or smaller.")
    return name, media_type, raw


def extract_pdf(name: str, raw: bytes) -> dict[str, object]:
    try:
        from pypdf import PdfReader
    except ImportError as error:
        raise RuntimeError("PDF support is not installed. Run: pip install -r requirements.txt") from error
    try:
        reader = PdfReader(BytesIO(raw))
        page_count = len(reader.pages)
        pages = []
        for number, page in enumerate(reader.pages[:MAX_PDF_PAGES], start=1):
            text = (page.extract_text() or "").strip()
            if text:
                pages.append(f"[Page {number}]\n{text}")
        extracted = "\n\n".join(pages)[:MAX_DOCUMENT_CHARS]
    except Exception as error:
        raise ValueError(f"{name}: the PDF could not be read. It may be encrypted or damaged.") from error

    images: list[str] = []
    # Scanned PDFs need visual page rendering; keep this small so the model request stays responsive.
    if len(extracted) < 80:
        try:
            import fitz

            document = fitz.open(stream=raw, filetype="pdf")
            for page in list(document)[:MAX_RENDERED_PAGES]:
                pixmap = page.get_pixmap(matrix=fitz.Matrix(1.25, 1.25), alpha=False)
                image_bytes = pixmap.tobytes("png")
                images.append("data:image/png;base64," + base64.b64encode(image_bytes).decode("ascii"))
            document.close()
        except ImportError as error:
            raise RuntimeError("Scanned PDF support is not installed. Run: pip install -r requirements.txt") from error
        except Exception as error:
            raise ValueError(f"{name}: scanned pages could not be rendered.") from error
    return {"name": name, "media_type": "application/pdf", "text": extracted, "images": images, "pages": page_count}


def prepare_attachments(items: object) -> list[dict[str, object]]:
    if items in (None, []):
        return []
    if not isinstance(items, list):
        raise ValueError("Attachments must be sent as a list.")
    if len(items) > MAX_ATTACHMENTS:
        raise ValueError(f"Attach no more than {MAX_ATTACHMENTS} files at a time.")
    prepared: list[dict[str, object]] = []
    total = 0
    for item in items:
        name, media_type, raw = decode_attachment(item)
        total += len(raw)
        if total > MAX_TOTAL_ATTACHMENT_BYTES:
            raise ValueError("The combined upload must be 20 MB or smaller.")
        if media_type == "application/pdf":
            prepared.append(extract_pdf(name, raw))
        else:
            prepared.append({
                "name": name,
                "media_type": media_type,
                "text": "",
                "images": [f"data:{media_type};base64," + base64.b64encode(raw).decode("ascii")],
                "pages": 1,
            })
    return prepared


def attachment_metadata(documents: list[dict[str, object]]) -> list[dict[str, object]]:
    return [
        {
            "name": item["name"],
            "type": item["media_type"],
            "pages": item["pages"],
            "textExtracted": bool(item["text"]),
        }
        for item in documents
    ]


def has_quiz_intent(question: str) -> bool:
    """Recognise explicit requests for quizzes/question sets in normal chat."""
    normalized = re.sub(r"\s+", " ", question.lower()).strip()
    return bool(re.search(
        r"\b(?:quiz(?:zes)?|mcqs?|multiple[ -]choice|questions?|question set|test me)\b",
        normalized,
    ))


def parse_quiz_response(answer: str) -> dict[str, object] | None:
    """Validate the small JSON contract used by the browser quiz renderer."""
    candidate = answer.strip()
    if candidate.startswith("```"):
        candidate = re.sub(r"^```(?:json)?\s*|\s*```$", "", candidate, flags=re.IGNORECASE)
    attempts = [candidate]
    start, end = candidate.find("{"), candidate.rfind("}")
    if start >= 0 and end > start:
        attempts.append(candidate[start : end + 1])
    for raw in attempts:
        try:
            payload = json.loads(raw)
        except (json.JSONDecodeError, TypeError):
            continue
        if not isinstance(payload, dict) or not isinstance(payload.get("questions"), list):
            continue
        questions: list[dict[str, object]] = []
        for item in payload["questions"][:7]:
            if not isinstance(item, dict):
                continue
            question = str(item.get("question", "")).strip()
            options = [str(option).strip() for option in item.get("options", []) if str(option).strip()]
            try:
                answer_index = int(item.get("answerIndex"))
            except (TypeError, ValueError):
                answer_index = -1
            if question and 2 <= len(options) <= 6 and 0 <= answer_index < len(options):
                questions.append({
                    "question": question,
                    "options": options,
                    "answerIndex": answer_index,
                    "explanation": str(item.get("explanation", "")).strip(),
                })
        if questions:
            return {
                "title": str(payload.get("title") or "Anusha's practice quiz").strip(),
                "context": str(payload.get("context") or payload.get("caseStudy") or payload.get("scenario") or "").strip(),
                "questions": questions,
            }
    return None


def generate_huggingface_answer(
    question: str,
    curriculum: list[dict[str, str]],
    sources: list[dict[str, str]],
    documents: list[dict[str, object]] | None = None,
    mode: str = "chat",
    route: dict[str, str] | None = None,
    history: list[dict[str, str]] | None = None,
) -> str:
    config, missing = huggingface_config()
    if not config:
        raise RuntimeError("Configure " + ", ".join(missing) + " locally before enabling live AI.")
    local_context = "\n\n".join(f"[{item['source']}]\n{item['excerpt']}" for item in curriculum) or "No directly matching local excerpt was found."
    web_context = "\n".join(f"- {item['title']}: {item['url']}" for item in sources) or "(None.)"
    documents = documents or []
    route = route or route_anusha_request(question, documents, mode, False)
    mode_instructions = {
        "chat": "Answer the user's question directly. When files are attached, ground the answer in them and name the relevant file.",
        "summarize": "Summarize the attached material with: executive overview, key points, supply-chain implications, actions, and important terms. Do not invent missing content.",
        "quiz": "Create exactly five multiple-choice learning questions that follow the user's requested subject, format and scenario. If the user asks for a case study, scenario or simulation, first write a concise but decision-rich fictional supply-chain case in the context field, then make every question depend on facts or trade-offs in that case; do not fall back to generic recall questions. If no scenario is requested, context may be an empty string. Ground uploaded-file quizzes in the uploaded material and curriculum quizzes in the supplied local excerpts. Return ONLY valid JSON, with no Markdown or prefacing text, in this exact shape: {\"title\":\"Short quiz title\",\"context\":\"Case study or empty string\",\"questions\":[{\"question\":\"Question text\",\"options\":[\"Option A\",\"Option B\",\"Option C\",\"Option D\"],\"answerIndex\":0,\"explanation\":\"Why this answer is correct using the case or source material\"}]}. answerIndex is zero-based. Use 2 to 6 plausible options per question. Never introduce question facts that are absent from the context or source material.",
    }
    mode = mode if mode in mode_instructions else "chat"
    system = f"""You are Anusha, a knowledgeable supply-chain planning, retail TPM and operations practice partner. You are operating as {route['name']}. {mode_instructions[mode]} Use your own expertise only to explain or contextualise; clearly distinguish it from facts found in the user's files. For images and scanned PDFs, work like a careful OCR analyst before answering: read all visible text at the highest available resolution, preserve headings, numbers, labels and table relationships, then answer the request from that transcription. If any region is unreadable, identify that region precisely instead of guessing. Never claim to have read content that is not visible or extracted."""
    document_context = []
    for document in documents:
        extracted = str(document.get("text") or "").strip()
        if extracted:
            document_context.append(f"[UPLOADED DOCUMENT: {document['name']}]\n{extracted}")
        else:
            document_context.append(f"[UPLOADED VISUAL: {document['name']}]\nInspect the corresponding image content in this message.")
    uploaded_context = "\n\n".join(document_context) or "(None.)"
    user_content = f"""USER QUESTION:
{question}

REQUEST MODE: {mode}

UPLOADED MATERIAL (ephemeral; not saved by the app):
{uploaded_context}

LOCAL CURRICULUM (private to this local app; never include saved notes):
{local_context}

WEB RESEARCH LEADS:
{web_context}"""
    user_message: str | list[dict[str, object]] = user_content
    visual_parts: list[dict[str, object]] = [{"type": "text", "text": user_content}]
    for document in documents:
        for image_url in document.get("images", []):
            visual_parts.append({"type": "image_url", "image_url": {"url": image_url}})
    if len(visual_parts) > 1:
        user_message = visual_parts
    has_visuals = len(visual_parts) > 1
    primary_model = config[route["model_key"]]
    candidate_models = [primary_model]
    if has_visuals or route["model_key"] == "multimodal_model_id":
        candidate_models.extend([
            config["multimodal_fallback_model_id"],
            "Qwen/Qwen2-VL-72B-Instruct",
        ])
    elif documents:
        # A text-extractable PDF does not require a vision model, so the
        # general Qwen route is a useful final fallback.
        candidate_models.append(config["general_model_id"])
    candidate_models = list(dict.fromkeys(model for model in candidate_models if model))
    conversation_messages: list[dict[str, object]] = [{"role": "system", "content": system}]
    conversation_messages.extend(history or [])
    conversation_messages.append({"role": "user", "content": user_message})

    payload = None
    provider_errors: list[str] = []
    for model_id in candidate_models:
        request_body = {
            "model": model_id,
            "messages": conversation_messages,
            "max_tokens": 1200 if mode in {"summarize", "quiz"} else 900,
            "temperature": 0.2,
        }
        chat_request = Request(
            "https://router.huggingface.co/v1/chat/completions",
            data=json.dumps(request_body).encode(),
            headers={"Authorization": "Bearer " + config["token"], "Content-Type": "application/json"},
            method="POST",
        )
        try:
            with trusted_urlopen(chat_request, timeout=None) as chat_response:
                payload = json.loads(chat_response.read().decode("utf-8"))
            break
        except HTTPError as error:
            body = error.read().decode("utf-8", errors="ignore")[:500]
            provider_errors.append(f"{model_id}: HTTP {error.code} {body}".strip())
            # Authentication and permission errors apply to every model and
            # should be fixed by the user rather than hidden by more retries.
            if error.code in {401, 403}:
                raise RuntimeError(f"Hugging Face authentication failed with HTTP {error.code}. Check HF_TOKEN and model access.") from error
            if error.code not in {400, 404, 405, 422, 424, 429, 500, 502, 503, 504}:
                raise RuntimeError(f"Hugging Face rejected {model_id} with HTTP {error.code}: {body}") from error
    if payload is None:
        details = " | ".join(provider_errors)
        raise RuntimeError(
            "No provider-backed Qwen vision route accepted this image request. "
            "Set HF_MULTIMODAL_MODEL_ID or HF_MULTIMODAL_FALLBACK_MODEL_ID to a Hugging Face Inference Providers vision model. "
            + details
        )
    answer = str(payload["choices"][0]["message"]["content"]).strip()
    # Quiz mode has a strict JSON contract. Do not trim it to the last sentence,
    # otherwise the closing braces and radio-button payload are lost.
    if mode != "quiz" and answer and answer[-1] not in ".!?`":
        complete = re.findall(r".*?[.!?](?:\s|$)", answer, flags=re.S)
        if complete:
            answer = "".join(complete).strip()
    return answer or "I could not form a complete response. Please try a more specific question."
class TradeFlowHandler(SimpleHTTPRequestHandler):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, directory=str(ROOT), **kwargs)

    def send_json(self, payload: dict, status: int = HTTPStatus.OK) -> None:
        body = json.dumps(payload).encode("utf-8")
        self.send_response(status)
        self.send_header("Content-Type", "application/json; charset=utf-8")
        self.send_header("Content-Length", str(len(body)))
        self.end_headers()
        self.wfile.write(body)

    def read_json(self) -> dict:
        size = int(self.headers.get("Content-Length", "0"))
        if size > MAX_REQUEST_BYTES:
            raise ValueError("Request is too large. Keep combined uploads below 20 MB.")
        return json.loads(self.rfile.read(size) or b"{}")

    def do_GET(self) -> None:
        request = urlparse(self.path)
        if request.path == "/":
            self.path = "/tradeflow-atlas.html"
            return super().do_GET()
        if request.path == "/api/state":
            zone = parse_qs(request.query).get("zone", ["Home"])[0]
            with connection() as db:
                capsule = db.execute("SELECT content, unlock_date FROM capsules WHERE zone = ?", (zone,)).fetchone()
                progress = db.execute("SELECT checked_keys FROM progress WHERE zone = ?", (zone,)).fetchone()
            return self.send_json({
                "note": "",
                "capsule": dict(capsule) if capsule else None,
                "checkedKeys": json.loads(progress["checked_keys"]) if progress else [],
            })
        if request.path == "/api/settings/theme":
            with connection() as db:
                row = db.execute("SELECT value FROM settings WHERE key = 'theme'").fetchone()
            return self.send_json({"theme": row["value"] if row else "light"})
        if request.path == "/api/assistant/health":
            config, missing = huggingface_config()
            return self.send_json({
                "ready": bool(config),
                "packageReady": True,
                "missingEnvironment": missing,
                "privacy": "Saved notes and time capsules are never sent to the AI or web search.",
            })
        if request.path == "/api/health":
            return self.send_json({
                "ok": True,
                "service": "TradeFlow Atlas",
                "newsConfigured": bool(news_api_key()),
            })
        if request.path == "/api/notes/archive":
            with connection() as db:
                entries = [
                    {"zone": x["zone"], "content": plain_note(x["content"]), "created_at": x["created_at"]}
                    for x in db.execute(
                        "SELECT zone, content, created_at FROM note_entries ORDER BY id DESC"
                    )
                ]
            return self.send_json({"entries": entries})
        if request.path == "/api/capsules/archive":
            with connection() as db:
                entries = [
                    {
                        "zone": x["zone"],
                        "content": x["content"],
                        "unlock_date": x["unlock_date"],
                        "updated_at": x["updated_at"],
                        "unlocked": x["unlock_date"] <= __import__("datetime").date.today().isoformat(),
                    }
                    for x in db.execute(
                        "SELECT zone, content, unlock_date, updated_at FROM capsules ORDER BY unlock_date"
                    )
                ]
            return self.send_json({"entries": entries})
        if request.path == "/api/industry-signals":
            force = parse_qs(request.query).get("refresh", [""])[0] == "1"
            return self.send_json(industry_signals(force=force))
        if request.path == "/api/export.xlsx":
            with connection() as db:
                payload = {
                    "notes": [{"zone": x["zone"], "content": plain_note(x["content"]), "updated_at": x["created_at"]} for x in db.execute("SELECT zone, content, created_at FROM note_entries ORDER BY id DESC")],
                    "capsules": [dict(x) for x in db.execute("SELECT zone, content, unlock_date, updated_at FROM capsules ORDER BY unlock_date")],
                    "progress": [{**dict(x), "checked_keys": progress_text(x["checked_keys"])} for x in db.execute("SELECT zone, checked_keys, updated_at FROM progress ORDER BY updated_at DESC")],
                    "settings": [dict(x) for x in db.execute("SELECT key, value, updated_at FROM settings ORDER BY key")],
                }
            export = build_excel_export(payload)
            self.send_response(HTTPStatus.OK)
            self.send_header("Content-Type", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
            self.send_header("Content-Disposition", "attachment; filename=tradeflow-atlas-export.xlsx")
            self.send_header("Content-Length", str(len(export)))
            self.end_headers()
            self.wfile.write(export)
            return
        return super().do_GET()

    def do_POST(self) -> None:
        request = urlparse(self.path)
        if request.path != "/api/assistant":
            return self.send_json({"error": "Unknown endpoint"}, HTTPStatus.NOT_FOUND)
        try:
            data = self.read_json()
        except json.JSONDecodeError:
            return self.send_json({"error": "Invalid JSON request"}, HTTPStatus.BAD_REQUEST)
        except ValueError as error:
            return self.send_json({"error": str(error)}, HTTPStatus.REQUEST_ENTITY_TOO_LARGE)
        mode = str(data.get("mode", "chat")).lower()
        if mode not in {"chat", "summarize", "quiz"}:
            return self.send_json({"error": "Unknown assistant mode."}, HTTPStatus.BAD_REQUEST)
        try:
            documents = prepare_attachments(data.get("attachments", []))
            history = prepare_conversation_history(data.get("history", []))
        except (ValueError, RuntimeError) as error:
            return self.send_json({"error": str(error)}, HTTPStatus.BAD_REQUEST)
        question = str(data.get("question", "")).strip()
        if not question and not documents:
            return self.send_json({"error": "Ask a question or attach a file first."}, HTTPStatus.BAD_REQUEST)
        if not question:
            question = "Summarize the attached material." if mode == "summarize" else "Create a quiz from the attached material." if mode == "quiz" else "Describe and extract the important information from the attached material."
        # A quiz request should produce the interactive quiz contract even when it
        # was typed into the normal Send box rather than submitted via Make quiz.
        if has_quiz_intent(question):
            mode = "quiz"
        previous_user_turn = next(
            (item["content"] for item in reversed(history) if item["role"] == "user"),
            "",
        )
        context_question = question
        if previous_user_turn:
            context_question = f"{previous_user_turn[-600:]}\nFollow-up: {question}"
        curriculum = relevant_curriculum(context_question)
        notion_context = notion_relevant_context(context_question, limit=2)
        if notion_context:
            curriculum = notion_context + curriculum
        # Document summaries and quizzes stay grounded in the upload unless the user asks a separate question.
        use_web = bool(data.get("useWeb", True)) and not (documents and mode in {"summarize", "quiz"})
        sources = web_search(context_question) if use_web else []
        route = route_anusha_request(question, documents, mode, use_web)
        try:
            answer = generate_huggingface_answer(question, curriculum, sources, documents, mode, route, history)
        except RuntimeError as error:
            if documents:
                return self.send_json({
                    "error": str(error),
                    "curriculum": curriculum,
                    "sources": sources,
                    "attachments": attachment_metadata(documents),
                    "route": route,
                }, HTTPStatus.BAD_GATEWAY)
            return self.send_json({
                "answer": local_guidance(question, curriculum),
                "notice": str(error),
                "curriculum": curriculum,
                "sources": sources,
                "mode": "local-guidance",
                "route": route,
            })
        except Exception as error:
            import traceback
            traceback.print_exc()
            return self.send_json({
                "error": f"The AI model could not complete that request: {error}",
                "curriculum": curriculum,
                "sources": sources,
                "attachments": attachment_metadata(documents),
                "route": route,
            }, HTTPStatus.BAD_GATEWAY)
        quiz = parse_quiz_response(answer) if mode == "quiz" else None
        return self.send_json({
            "answer": "" if quiz else answer,
            "quiz": quiz,
            "curriculum": curriculum,
            "sources": sources,
            "attachments": attachment_metadata(documents),
            "mode": mode,
            "route": route,
        })

    def do_PUT(self) -> None:
        request = urlparse(self.path)
        data = self.read_json()
        with connection() as db:
            if request.path == "/api/notes":
                db.execute("INSERT INTO note_entries(zone, content) VALUES(?, ?)", (data["zone"], data.get("content", "")))
            elif request.path == "/api/capsules":
                db.execute("INSERT INTO capsules(zone, content, unlock_date) VALUES(?, ?, ?) ON CONFLICT(zone) DO UPDATE SET content=excluded.content, unlock_date=excluded.unlock_date, updated_at=CURRENT_TIMESTAMP", (data["zone"], data["content"], data["unlockDate"]))
            elif request.path == "/api/progress":
                db.execute("INSERT INTO progress(zone, checked_keys) VALUES(?, ?) ON CONFLICT(zone) DO UPDATE SET checked_keys=excluded.checked_keys, updated_at=CURRENT_TIMESTAMP", (data["zone"], json.dumps(data.get("checkedKeys", []))))
            elif request.path == "/api/settings/theme":
                db.execute("INSERT INTO settings(key, value) VALUES('theme', ?) ON CONFLICT(key) DO UPDATE SET value=excluded.value, updated_at=CURRENT_TIMESTAMP", (data["theme"],))
            else:
                return self.send_json({"error": "Unknown endpoint"}, HTTPStatus.NOT_FOUND)
        self.send_json({"ok": True})


if __name__ == "__main__":
    host = os.environ.get("TRADEFLOW_HOST", "0.0.0.0")
    port = int(os.environ.get("PORT", os.environ.get("TRADEFLOW_PORT", "8799")))
    print(f"TradeFlow Atlas is running at http://{host}:{port}")
    ThreadingHTTPServer((host, port), TradeFlowHandler).serve_forever()
